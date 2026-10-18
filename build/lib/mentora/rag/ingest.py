"""Ingest course material (markdown, text, PDF, slides, FAQ) into ~300-token passages
tagged with concepts, section and week."""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path

import yaml
from sqlalchemy import delete
from sqlalchemy.orm import Session, sessionmaker

from mentora.course import ConceptSpec, CourseSpec
from mentora.db import models as m
from mentora.rag.embeddings import Embedder
from mentora.rag.text import count_tokens, slug, split_sentences

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)
TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
FAQ_SUFFIXES = {".yaml", ".yml"}


@dataclass(frozen=True)
class Section:
    title: str | None
    text: str


@dataclass(frozen=True)
class Document:
    source: str  # path relative to the corpus root
    title: str
    week: int | None
    sections: tuple[Section, ...]


@dataclass(frozen=True)
class PassageDraft:
    id: str
    text: str
    source_file: str
    section: str | None
    week: int | None
    concept_keys: tuple[str, ...]
    token_count: int
    embedding: tuple[float, ...] | None = field(default=None, compare=False)


# --- loading ---------------------------------------------------------------------------


def _front_matter(raw: str) -> tuple[dict[str, object], str]:
    if raw.startswith("---"):
        parts = raw.split("---", 2)
        if len(parts) == 3:
            return yaml.safe_load(parts[1]) or {}, parts[2]
    return {}, raw


def _markdown_sections(body: str) -> list[Section]:
    sections: list[Section] = []
    matches = list(_HEADING_RE.finditer(body))
    if not matches or matches[0].start() > 0:
        head = body[: matches[0].start()] if matches else body
        if head.strip():
            sections.append(Section(None, head))
    for i, mt in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
        text = body[mt.end() : end]
        if text.strip():
            sections.append(Section(mt.group(2).strip(), text))
    return sections


def _load_text(path: Path, rel: str) -> Document:
    meta, body = _front_matter(path.read_text(encoding="utf-8"))
    week = meta.get("week")
    return Document(
        source=rel,
        title=str(meta.get("title", path.stem)),
        week=int(week) if isinstance(week, int | str) and str(week).isdigit() else None,
        sections=tuple(_markdown_sections(body)),
    )


def _load_pdf(path: Path, rel: str, week: int | None) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections = tuple(
        Section(f"Page {i}", page.extract_text() or "")
        for i, page in enumerate(reader.pages, start=1)
    )
    return Document(rel, path.stem, week, tuple(s for s in sections if s.text.strip()))


def _load_pptx(path: Path, rel: str, week: int | None) -> Document:
    from pptx import Presentation

    sections = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape is not slide.shapes.title
        ]
        if title or any(t.strip() for t in texts):
            sections.append(Section(f"Slide {i}: {title}".rstrip(": "), "\n\n".join(texts)))
    return Document(rel, path.stem, week, tuple(sections))


def _load_faq(path: Path, rel: str) -> list[Document]:
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    docs = []
    for i, qa in enumerate(raw.get("faq", []), start=1):
        week = qa.get("week")
        docs.append(
            Document(
                source=f"{rel}#q{i}",
                title=f"FAQ: {qa['q']}",
                week=int(week) if week is not None else None,
                sections=(Section("FAQ", f"Q: {qa['q']}\n\nA: {qa['a']}"),),
            )
        )
    return docs


def _week_from_name(path: Path) -> int | None:
    match = re.search(r"week[_\- ]?(\d+)", path.stem, re.IGNORECASE)
    return int(match.group(1)) if match else None


def load_documents(root: Path) -> list[Document]:
    """Load every supported file under ``root`` (sorted, so ids are stable)."""
    files = [root] if root.is_file() else sorted(p for p in root.rglob("*") if p.is_file())
    base = root.parent if root.is_file() else root
    docs: list[Document] = []
    for path in files:
        rel = path.relative_to(base).as_posix()
        suffix = path.suffix.lower()
        if suffix in TEXT_SUFFIXES:
            doc = _load_text(path, rel)
            docs.append(doc if doc.week is not None else replace(doc, week=_week_from_name(path)))
        elif suffix in FAQ_SUFFIXES:
            docs.extend(_load_faq(path, rel))
        elif suffix == ".pdf":
            docs.append(_load_pdf(path, rel, _week_from_name(path)))
        elif suffix == ".pptx":
            docs.append(_load_pptx(path, rel, _week_from_name(path)))
    return docs


# --- chunking and tagging --------------------------------------------------------------


def _units(text: str, max_tokens: int) -> list[str]:
    """Paragraphs, with over-long paragraphs split into sentences."""
    units: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = " ".join(para.split())
        if not para:
            continue
        if count_tokens(para) <= max_tokens:
            units.append(para)
        else:
            units.extend(split_sentences(para))
    return units


def chunk_section(text: str, chunk_tokens: int, overlap_tokens: int) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    size = 0
    for unit in _units(text, chunk_tokens):
        n = count_tokens(unit)
        if current and size + n > chunk_tokens:
            chunks.append(" ".join(current))
            # Carry trailing sentences of the previous chunk as overlap.
            tail: list[str] = []
            for sent in reversed(split_sentences(" ".join(current))):
                if count_tokens(" ".join([sent, *tail])) > overlap_tokens:
                    break
                tail.insert(0, sent)
            current, size = tail, count_tokens(" ".join(tail))
        current.append(unit)
        size += n
    if current:
        chunks.append(" ".join(current))
    return chunks


def _singular(word: str) -> str:
    if word.endswith("ies"):
        return word[:-3] + "y"
    if word.endswith(("sses", "xes")):
        return word[:-2]
    if word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


def _phrase_pattern(phrase: str) -> re.Pattern[str]:
    """Match a phrase in singular or plural form (only the last word is inflected)."""
    *head, last = phrase.lower().split()
    stem = _singular(last)
    if stem.endswith("y"):
        tail = rf"{re.escape(stem[:-1])}(?:y|ies)"
    else:
        tail = rf"{re.escape(stem)}(?:s|es)?"
    prefix = "".join(rf"{re.escape(w)}\s+" for w in head)
    return re.compile(rf"\b{prefix}{tail}\b")


def tag_concepts(text: str, concepts: Sequence[ConceptSpec]) -> tuple[str, ...]:
    lowered = text.lower()
    return tuple(
        c.key for c in concepts if any(_phrase_pattern(p).search(lowered) for p in c.phrases)
    )


# Text in course material that tries to instruct the tutor (prompt injection via files).
INJECTION_RE = re.compile(
    r"\b(ignore|disregard|forget|override)\b[^.\n]{0,40}\b(rules|instructions|guidelines|"
    r"previous|above|system)\b|\bsystem prompt\b|\byou are now\b|\bact as (an?|the)\b|"
    r"\b(give|tell|show|reveal)\b[^.\n]{0,30}\b(students|learners|the learner|everyone)\b"
    r"[^.\n]{0,30}\b(the |all |full |complete |correct )+(answers?|solutions?)\b|"
    r"\bdo not tell the teacher\b",
    re.IGNORECASE,
)


def is_suspicious(text: str) -> bool:
    return bool(INJECTION_RE.search(text))


def build_passages(
    docs: Iterable[Document],
    course: CourseSpec,
    chunk_tokens: int,
    overlap_tokens: int,
    quarantine: list[PassageDraft] | None = None,
) -> list[PassageDraft]:
    """Chunk documents into passages. Chunks that look like instructions to the tutor are
    never stored; they are collected in ``quarantine`` for a human to review."""
    out: list[PassageDraft] = []
    prefix = slug(course.code, 12)
    for doc in docs:
        path, _, anchor = doc.source.partition("#")
        stem = slug(path.rsplit(".", 1)[0] + (f"-{anchor}" if anchor else ""), 36)
        n = 0
        for section in doc.sections:
            for chunk in chunk_section(section.text, chunk_tokens, overlap_tokens):
                n += 1
                heading = section.title or doc.title
                target = quarantine if is_suspicious(chunk) else out
                if target is None:
                    continue
                target.append(
                    PassageDraft(
                        id=f"{prefix}-{stem}-p{n:03d}",
                        text=chunk,
                        source_file=doc.source,
                        section=heading,
                        week=doc.week,
                        concept_keys=tag_concepts(f"{heading} {chunk}", course.concepts),
                        token_count=count_tokens(chunk),
                    )
                )
    return out


def embed_passages(passages: Sequence[PassageDraft], embedder: Embedder) -> list[PassageDraft]:
    vectors = embedder.embed([f"{p.section or ''}\n{p.text}" for p in passages])
    return [replace(p, embedding=tuple(v)) for p, v in zip(passages, vectors, strict=True)]


def store_passages(
    session_factory: sessionmaker[Session], course_id: int, passages: Sequence[PassageDraft]
) -> int:
    """Replace this course's passages from the same source files, then insert."""
    sources = {p.source_file.split("#", 1)[0] for p in passages}
    with session_factory.begin() as db:
        for src in sources:
            db.execute(
                delete(m.CorpusPassage).where(
                    m.CorpusPassage.course_id == course_id,
                    (m.CorpusPassage.source_file == src)
                    | m.CorpusPassage.source_file.startswith(f"{src}#"),
                )
            )
        db.add_all(
            m.CorpusPassage(
                id=p.id,
                course_id=course_id,
                text=p.text,
                embedding=list(p.embedding) if p.embedding is not None else None,
                concept_keys=list(p.concept_keys),
                section=p.section,
                week=p.week,
                source_file=p.source_file,
                token_count=p.token_count,
            )
            for p in passages
        )
    return len(passages)
