from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from mentora.course import CourseSpec, upsert_course
from mentora.db import StoreFactory
from mentora.db import models as m
from mentora.rag.ingest import (
    PassageDraft,
    build_passages,
    chunk_section,
    load_documents,
    store_passages,
    tag_concepts,
)
from mentora.rag.text import count_tokens

ROOT = Path(__file__).resolve().parents[3]
CORPUS = ROOT / "data" / "sample" / "corpus"


def test_sample_corpus_passages(passages: list[PassageDraft]) -> None:
    ids = [p.id for p in passages]
    assert len(ids) == len(set(ids)), "passage ids must be unique"
    assert all(len(i) <= 64 for i in ids)
    assert all(p.token_count <= 300 + 60 for p in passages)
    assert all(p.embedding is not None and len(p.embedding) == 384 for p in passages)
    by_source = {p.source_file.split("#")[0] for p in passages}
    assert by_source == {
        "week1_variables.md",
        "week2_conditionals.md",
        "week3_loops.md",
        "week4_lists_tuples.md",
        "week5_dicts_sets.md",
        "week6_functions.md",
        "week7_recursion.md",
        "week8_files_exceptions.md",
        "week9_classes.md",
        "week10_complexity.md",
        "faq.yaml",
    }


def test_front_matter_week_and_sections(passages: list[PassageDraft]) -> None:
    loops = [p for p in passages if p.source_file == "week3_loops.md"]
    assert {p.week for p in loops} == {3}
    assert {"While loops", "For loops and ranges", "Strings"} <= {p.section for p in loops}


def test_concept_tagging(passages: list[PassageDraft], course: CourseSpec) -> None:
    for_loop = next(p for p in passages if p.section == "For loops and ranges")
    assert "for_loops" in for_loop.concept_keys
    assert tag_concepts("Recursion means a function calls itself", course.concepts) == (
        "functions",
        "recursion",
    )


def test_faq_entries_become_their_own_passages(passages: list[PassageDraft]) -> None:
    faq = [p for p in passages if p.source_file.startswith("faq.yaml#")]
    assert len(faq) == 3
    assert faq[0].text.startswith("Q: Why does input()")


def test_ids_are_stable_across_runs(course: CourseSpec) -> None:
    first = [p.id for p in build_passages(load_documents(CORPUS), course, 300, 40)]
    second = [p.id for p in build_passages(load_documents(CORPUS), course, 300, 40)]
    assert first == second


def test_long_sections_are_chunked_with_overlap() -> None:
    text = " ".join(f"Sentence number {i} talks about loops." for i in range(120))
    chunks = chunk_section(text, chunk_tokens=100, overlap_tokens=20)
    assert len(chunks) > 5
    assert all(count_tokens(c) <= 100 + 20 for c in chunks)
    first_tail = chunks[0].split(". ")[-1].rstrip(".")
    assert first_tail in chunks[1]  # overlap carries the last sentence forward


def test_pptx_slides(tmp_path: Path, course: CourseSpec) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Dictionaries"
    slide.placeholders[1].text = "A dictionary maps keys to values."
    path = tmp_path / "week5_dicts.pptx"
    deck.save(str(path))
    [doc] = load_documents(path)
    assert doc.week == 5
    [p] = build_passages([doc], course, 300, 40)
    assert p.section == "Slide 1: Dictionaries"
    assert "maps keys to values" in p.text
    assert "dicts" in p.concept_keys


def _tiny_pdf(text: str) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def test_pdf_pages(tmp_path: Path, course: CourseSpec) -> None:
    path = tmp_path / "week8_exceptions.pdf"
    path.write_bytes(_tiny_pdf("Exceptions signal errors at run time."))
    [doc] = load_documents(path)
    [p] = build_passages([doc], course, 300, 40)
    assert p.section == "Page 1" and p.week == 8
    assert "Exceptions signal errors" in p.text
    assert "exceptions" in p.concept_keys


def test_store_replaces_passages_from_same_source(
    session_factory: StoreFactory, course: CourseSpec, passages: list[PassageDraft]
) -> None:
    with session_factory.begin() as db:
        course_id = upsert_course(db, course).id
    store_passages(session_factory, course_id, passages)
    store_passages(session_factory, course_id, passages)  # re-ingest: no duplicates
    with session_factory() as db:
        n = db.count(m.CorpusPassage)
        stored = db.get(m.CorpusPassage, passages[0].id)
    assert n == len(passages)
    assert stored is not None and stored.embedding is not None
    assert stored.embedding == pytest.approx(list(passages[0].embedding or ()))


def test_upsert_course_creates_prerequisite_edges(
    session_factory: StoreFactory, course: CourseSpec
) -> None:
    with session_factory.begin() as db:
        upsert_course(db, course)
        upsert_course(db, course)  # idempotent
    with session_factory() as db:
        assert db.count(m.Concept) == len(course.concepts)
        edges = [p for c in db.find(m.Concept) for p in c.prerequisite_ids]
    assert len(edges) == sum(len(c.prereqs) for c in course.concepts)


def test_cyclic_prerequisites_rejected() -> None:
    with pytest.raises(ValueError, match="cycle"):
        CourseSpec(
            code="X",
            title="x",
            concepts=[
                {"key": "a", "name": "A", "prereqs": ["b"]},  # type: ignore[list-item]
                {"key": "b", "name": "B", "prereqs": ["a"]},  # type: ignore[list-item]
            ],
        )
